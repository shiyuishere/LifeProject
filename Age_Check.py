import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Set file paths
input_file = "/Users/fancy/Documents/Utrecht SaSR/SoDa/FTOLP/Data/life-project-questionnaire-main NL final/Analysis/Unified_Data.xlsx"
output_file = "/Users/fancy/Documents/Utrecht SaSR/SoDa/FTOLP/Data/life-project-questionnaire-main NL final/Analysis/Unified_Data_AgeChecked.xlsx"

# Read the input Excel file
df = pd.read_excel(input_file)

# Age column and goal age columns
age_col = df.columns[1]
goal_age_cols = [col for col in df.columns if "_age" in col]

# remarker for errors
for col in goal_age_cols:
    main_age = pd.to_numeric(df[age_col], errors="coerce")
    goal_age = pd.to_numeric(df[col], errors="coerce")
    error_mask = goal_age < main_age
    df.loc[error_mask, col] = "*" + df.loc[error_mask, col].astype(str)

# Save the modified DataFrame to a new Excel file
df.to_excel(output_file, index=False)

# Highlight cells with errors
wb = load_workbook(output_file)
ws = wb.active

blue_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")

for col_name in goal_age_cols:
    col_idx = df.columns.get_loc(col_name) + 1  # openpyxl starts counting from 1
    for row in range(2, ws.max_row + 1):  # start from the second row, skipping header
        cell = ws.cell(row=row, column=col_idx)
        if isinstance(cell.value, str) and cell.value.startswith("*"):
            cell.fill = blue_fill

# Save the workbook with highlighted errors
wb.save(output_file)
print("✅ Age checking complete. File saved to:", output_file)
