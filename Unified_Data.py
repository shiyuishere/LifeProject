import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Set file paths
file1 = "/Users/fancy/Documents/Utrecht SaSR/SoDa/FTOLP/Data/life-project-questionnaire-main NL final/Analysis/Coder1/Data_Coder1.xlsx"
file2 = "/Users/fancy/Documents/Utrecht SaSR/SoDa/FTOLP/Data/life-project-questionnaire-main NL final/Analysis/Coder2/Data_Coder2.xlsx"
output_file = "/Users/fancy/Documents/Utrecht SaSR/SoDa/FTOLP/Data/life-project-questionnaire-main NL final/Analysis/Unified_Data.xlsx"

# Load Excel files
df1 = pd.read_excel(file1)
df2 = pd.read_excel(file2)

# Identify category columns
category_cols = [col for col in df1.columns if "_category" in col]
id_col = df1.columns[0]  # Assuming first column is ID

# Make a copy for merged output
merged = df1.copy()

# Merge each category column
for col in category_cols:
    # Compare coder1 and coder2
    val1 = df1[col].fillna("").astype(str)
    val2 = df2[col].fillna("").astype(str)

    same = val1 == val2
    merged[col] = val1 + "/" + val2
    merged.loc[same, col] = val1[same]  # Keep coder1 result when same
    merged.loc[~same, col] = "*" + merged.loc[~same, col]  # Add * when different

# Mark IDs where any category is different
diff_any = (df1[category_cols].fillna("").astype(str) != df2[category_cols].fillna("").astype(str)).any(axis=1)
merged.loc[diff_any, id_col] = "*" + merged.loc[diff_any, id_col].astype(str)

# Save to Excel
merged.to_excel(output_file, index=False)

# Red cell fill
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

# Load the saved Excel file
wb = load_workbook(output_file)
ws = wb.active

# Get all categorical column indices (openpyxl starts counting from 1)
category_col_indices = [merged.columns.get_loc(col) + 1 for col in category_cols]

# Iterate over each target column and each row
for col_idx in category_col_indices:
    for row in range(2, ws.max_row + 1):  # Start from the second row, skipping header
        cell = ws.cell(row=row, column=col_idx)
        if isinstance(cell.value, str) and cell.value.startswith("*"):
            cell.fill = red_fill

# Save the workbook with highlighted errors
wb.save(output_file)
print("✅ Unified data saved to:", output_file)

